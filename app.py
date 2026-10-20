import os
import zipfile
import re
import uuid
import shutil
from pathlib import Path
from datetime import datetime
from flask import Flask, render_template, request, send_file, jsonify, send_from_directory
from werkzeug.utils import secure_filename
from PIL import Image, ImageEnhance, ImageFilter
import pytesseract
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 500 * 1024 * 1024
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['RESULTS_FOLDER'] = 'results'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['RESULTS_FOLDER'], exist_ok=True)

ALLOWED_EXTENSIONS = {'zip'}
ALLOWED_IMAGE_EXT = {'.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.webp'}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def preprocess_image(img):
    """Предобработка для лучшего OCR"""
    enhancer = ImageEnhance.Contrast(img)
    img = enhancer.enhance(2.0)
    img = img.filter(ImageFilter.SHARPEN)
    img = img.convert('L')
    width, height = img.size
    img = img.resize((width * 2, height * 2), Image.Resampling.LANCZOS)
    return img


def create_thumbnail(img_path, output_path, size=(300, 300)):
    """Создаёт миниатюру для предпросмотра"""
    try:
        with Image.open(img_path) as img:
            img.thumbnail(size, Image.Resampling.LANCZOS)
            img.save(output_path, 'JPEG', quality=80)
        return True
    except Exception:
        return False


def analyze_image(image_path):
    """Анализ изображения на наличие адреса"""
    try:
        img = Image.open(image_path)
        img_processed = preprocess_image(img)
        
        custom_config = r'--oem 3 --psm 6 -l rus+eng'
        text = pytesseract.image_to_string(img_processed, config=custom_config)
        text_lower = text.lower()
        
        patterns = {
            'улица': [r'ул\.\s*[\w\-]+', r'улица\s+[\w\-]+'],
            'проспект': [r'пр\.\s*[\w\-]+', r'проспект\s+[\w\-]+'],
            'переулок': [r'пер\.\s*[\w\-]+', r'переулок\s+[\w\-]+'],
            'дом': [r'д\.\s*\d+[\w\-/]*', r'дом\s+\d+[\w\-/]*'],
        }
        
        found = []
        for category, pats in patterns.items():
            for pattern in pats:
                matches = re.findall(pattern, text_lower)
                if matches:
                    found.extend(matches)
        
        if found:
            unique_found = list(dict.fromkeys(found))
            return True, f"Найдено: {', '.join(unique_found[:5])}"
        else:
            return False, "Адрес не распознан или недостаточно чёткий"
            
    except Exception as e:
        return False, f"Ошибка обработки: {str(e)}"


def make_excel_pretty(excel_path):
    """Делает Excel красивым: цвета, автоширина, автофильтр"""
    wb = load_workbook(excel_path)
    ws = wb.active
    ws.title = "Результаты анализа"
    
    # Стили
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    yes_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    yes_font = Font(color='006100', bold=True)
    
    no_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    no_font = Font(color='9C0006', bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    
    # Стилизуем заголовки
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    
    # Стилизуем строки данных
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        
        status_cell = row[1]  # колонка "Соответствует"
        if status_cell.value == 'Да':
            status_cell.fill = yes_fill
            status_cell.font = yes_font
        elif status_cell.value == 'Нет':
            status_cell.fill = no_fill
            status_cell.font = no_font
    
    # Автоширина колонок
    for col_idx, column in enumerate(ws.columns, 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for cell in column:
            try:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    max_line = max(len(line) for line in lines)
                    if max_line > max_length:
                        max_length = max_line
            except Exception:
                pass
        
        # Ограничиваем ширину
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[col_letter].width = max(adjusted_width, 15)
    
    # Фиксируем заголовок
    ws.freeze_panes = 'A2'
    
    # Автофильтр
    ws.auto_filter.ref = ws.dimensions
    
    # Высота заголовка
    ws.row_dimensions[1].height = 25
    
    wb.save(excel_path)


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'Нет файла'}), 400
    
    file = request.files['file']
    if file.filename == '' or not allowed_file(file.filename):
        return jsonify({'error': 'Разрешены только ZIP-архивы'}), 400
    
    task_id = str(uuid.uuid4())[:8]
    filename = secure_filename(file.filename)
    zip_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{task_id}_{filename}")
    file.save(zip_path)
    
    extract_dir = os.path.join(app.config['UPLOAD_FOLDER'], task_id)
    previews_dir = os.path.join(app.config['RESULTS_FOLDER'], task_id, 'previews')
    os.makedirs(previews_dir, exist_ok=True)
    
    try:
        with zipfile.ZipFile(zip_path, 'r') as zip_ref:
            zip_ref.extractall(extract_dir)
        
        # Собираем изображения
        image_files = []
        for root, dirs, files in os.walk(extract_dir):
            for f in files:
                if Path(f).suffix.lower() in ALLOWED_IMAGE_EXT:
                    image_files.append(os.path.join(root, f))
        
        if not image_files:
            return jsonify({'error': 'В архиве не найдено изображений'}), 400
        
        # Анализ
        results = []
        for img_path in image_files:
            img_filename = Path(img_path).name
            is_valid, comment = analyze_image(img_path)
            
            # Создаём превью
            preview_name = f"{Path(img_filename).stem}.jpg"
            preview_path = os.path.join(previews_dir, preview_name)
            create_thumbnail(img_path, preview_path)
            
            results.append({
                'filename': img_filename,
                'preview': preview_name,
                'valid': is_valid,
                'comment': comment
            })
        
        # Excel
        df = pd.DataFrame([
            {
                'Название файла': r['filename'],
                'Соответствует': 'Да' if r['valid'] else 'Нет',
                'Комментарий': r['comment']
            }
            for r in results
        ])
        
        output_filename = f"Результат_{Path(filename).stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(app.config['RESULTS_FOLDER'], f"{task_id}_{output_filename}")
        df.to_excel(output_path, index=False, engine='openpyxl')
        make_excel_pretty(output_path)
        
        valid_count = sum(1 for r in results if r['valid'])
        
        return jsonify({
            'success': True,
            'task_id': task_id,
            'total': len(results),
            'valid': valid_count,
            'invalid': len(results) - valid_count,
            'download_url': f'/download/{task_id}_{output_filename}',
            'results': results
        })
        
    except Exception as e:
        return jsonify({'error': f'Ошибка: {str(e)}'}), 500
    finally:
        if os.path.exists(zip_path):
            os.remove(zip_path)
        if os.path.exists(extract_dir):
            shutil.rmtree(extract_dir, ignore_errors=True)


@app.route('/download/<filename>')
def download_file(filename):
    file_path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    if os.path.exists(file_path):
        return send_file(file_path, as_attachment=True)
    return jsonify({'error': 'Файл не найден'}), 404


@app.route('/preview/<task_id>/<filename>')
def preview_file(task_id, filename):
    previews_dir = os.path.join(app.config['RESULTS_FOLDER'], task_id, 'previews')
    return send_from_directory(previews_dir, filename)


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=False)